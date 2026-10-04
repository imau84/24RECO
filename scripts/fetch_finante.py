#!/usr/bin/env python3
"""
Finanțe — 20 de statistici despre industria financiară din România → src/data/finante/date.json

Pachetul „Financiara” (scripts/financiara/, generat separat și copiat aici neschimbat) descarcă datele
din Eurostat, BCE, BIS, BVB (buletine PDF), ASF și EIOPA și scrie câte un Excel per set în
scripts/financiara/out/. Acest script:
  1. rulează cele 20 de scripturi scripts/financiara/scripts/financiara_NN_MM.py (fiecare separat:
     dacă unul eșuează, celelalte merg mai departe);
  2. citește din fiecare Excel foile de serii alese mai jos (rândul 3 = antet „Luna”/„Trimestru” + serii)
     și le transformă în formatul EurostatSet folosit de pagină;
  3. pentru seturile care n-au putut fi regenerate păstrează versiunea veche din JSON.

Rulează lunar din GitHub Actions (.github/workflows/update-finante.yml).
Rulare: pip install requests pandas numpy openpyxl ; apt install poppler-utils (pdftotext, pentru BVB)
        python scripts/fetch_finante.py              # descarcă + convertește
        python scripts/fetch_finante.py --doar-conversie [dir_xlsx]   # doar convertește Excel-urile existente
"""
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl

PKG = Path(__file__).resolve().parent / "financiara"
XLSX_DIR = PKG / "out"
OUT = Path("src/data/finante/date.json")

# paranteza de la sfârșitul unui antet e unitate doar dacă arată a unitate („Populație (gospodării)” nu e)
RE_UNIT = re.compile(r"\s*\(((?:mil|mld|mii)\.?[^()]*|%|lei|RON|EUR|puncte|persoane|p\.p\.)\)\s*$")
# seriile-flux (sume pe lună/trimestru) se adună pe an; restul (solduri, rate, indici) se mediază.
# Se testează pe antetul original din Excel; foile întregi de fluxuri au opțiunea "suma": True.
RE_SUMA = re.compile(r"^(Valoare tranzacționată|Nr\. tranzacții|Contribuții brute|Total prime brute|Prime brute|Daune brute|Cheltuieli de exploatare|.*flux net lunar)")

# Foile folosite din fiecare Excel. sheet: (nume foaie, coloane | None = toate, opțiuni)
#   unitate / zecimale: pentru coloanele fără unitate în antet;  nume: redenumiri;  prefix: „Activ net — ”
#   fara: coloane excluse când coloane=None
SETURI = {
    "piata-capital": {
        "label": "Piața de capital", "icon": "📈",
        "intro": "Bursa de la București: cât valorează companiile listate, cât se tranzacționează, indicii BET și dobânzile la care se împrumută statul și băncile.",
        "seturi": [
            {"cod": "01 01", "key": "randament", "scurt": "Dobânda la datoria statului", "icon": "🏛️",
             "descriere": "Cu ce dobândă se împrumută statul român pe 10 ani, comparat cu țările vecine și zona euro. "
                          "Cu cât e mai mare, cu atât datoria publică e mai scumpă — și, de regulă, și creditele pentru toată lumea.",
             "unitate": "% pe an", "zecimale": 2, "rata": True,
             "foi": [("Serii lunare", None, {})]},
            {"cod": "01 02", "key": "bvb", "scurt": "Bursa: capitalizare și tranzacții", "icon": "🏦",
             "descriere": "Cât valorează toate companiile listate la Bursa de Valori București (capitalizarea) și câți bani se tranzacționează lunar.",
             "zecimale": 1,
             "foi": [("Serii lunare", ["Capitalizare (mil. RON)", "Capitalizare (mil. EUR)", "Valoare tranzacționată (mil. RON)",
                                       "Valoare tranzacționată (mil. EUR)", "Valoare medie zilnică (mil. RON)", "Nr. tranzacții"],
                      {"unitate": "tranzacții", "zecimale": 0}),
                     ("Lichiditate", ["Rata de lichiditate anualizată (%)"], {"nume": {"Rata de lichiditate anualizată": "Lichiditate (cât din bursă se tranzacționează într-un an)"}})]},
            {"cod": "01 03", "key": "indici", "scurt": "Indicii bursei (BET)", "icon": "📊",
             "descriere": "Indicii Bursei de Valori București la sfârșitul fiecărei luni. BET urmărește cele mai mari companii; "
                          "BET-TR include și dividendele reinvestite, deci arată câștigul total al unui investitor.",
             "unitate": "puncte", "zecimale": 1,
             "foi": [("Serii lunare", None, {})]},
            {"cod": "01 04", "key": "robor", "scurt": "ROBOR și dobânzile interbancare", "icon": "💱",
             "descriere": "Dobânzile la care se împrumută băncile între ele (ROBOR), pe diferite termene. "
                          "ROBOR 3 luni e cel de care depind multe credite vechi în lei.",
             "unitate": "% pe an", "zecimale": 2, "rata": True,
             "foi": [("Maturități România", None, {"prefix": "ROBOR "}),
                     ("Comparație regională", ["Polonia", "Ungaria", "Cehia", "Zona euro"], {"prefix": "3 luni — "})]},
            {"cod": "01 05", "key": "sectoare", "scurt": "Bursa pe sectoare", "icon": "🏭",
             "descriere": "Cât valorează companiile listate la bursă, pe sectoare: bănci, energie, petrol și gaze, industrie etc.",
             "unitate": "mld. lei", "zecimale": 2,
             "foi": [("Capitalizare pe sectoare", None, {"nume": {"TOTAL": "Toată bursa"}})]},
        ],
    },
    "banci": {
        "label": "Bănci", "icon": "🏦",
        "intro": "Ce dobânzi plătim la credite și ce primim la depozite, câți bani au românii în bănci, cât s-au împrumutat și cât costă euro.",
        "seturi": [
            {"cod": "02 01", "key": "credite-dobanzi", "scurt": "Dobânzi la credite", "icon": "🏠",
             "descriere": "Dobânda medie la creditele noi în lei: pentru casă, de consum, descoperit de cont, pentru firme. "
                          "DAE = costul total, cu tot cu comisioane.",
             "unitate": "% pe an", "zecimale": 2, "rata": True,
             "foi": [("Rate lunare", None, {}), ("DAE si costul total", None, {})]},
            {"cod": "02 02", "key": "depozite-dobanzi", "scurt": "Dobânzi la depozite", "icon": "💰",
             "descriere": "Ce dobândă primești la un depozit nou sau în contul curent, pe termene, și cât câștigă banca din diferența dintre "
                          "dobânda la credite și cea la depozite (marja).",
             "unitate": "% pe an", "zecimale": 2, "rata": True,
             "foi": [("Dobanzi depozite noi", None, {}),
                     ("Marja bancara", ["Marjă credit locuință – depozit populație (p.p.)", "Marjă credit consum – depozit populație (p.p.)",
                                        "Marjă credit firme – depozit firme (p.p.)"], {})]},
            {"cod": "02 03", "key": "bilant", "scurt": "Credite și depozite", "icon": "📒",
             "descriere": "Câți bani au românii și firmele în bănci și cât datorează, la sfârșitul fiecărei luni, în miliarde de euro.",
             "unitate": "mld. euro", "zecimale": 2,
             "foi": [("Solduri lunare", None, {}),
                     ("Credite vs depozite", ["Raport credite / depozite (%)"], {}),
                     ("Ritm si valuta", ["Ponderea creditelor în valută (%)"], {})]},
            {"cod": "02 04", "key": "curs", "scurt": "Cursul leului și dobânda BNR", "icon": "💶",
             "descriere": "Câți lei costă un euro, un dolar, un franc elvețian și o liră (media lunii), plus dobânda de politică monetară a BNR.",
             "unitate": "lei", "zecimale": 4,
             "foi": [("Cursuri lunare", None, {"nume": {"RON/EUR": "1 euro", "RON/USD": "1 dolar american", "RON/CHF": "1 franc elvețian", "RON/GBP": "1 liră sterlină"}}),
                     ("Dobanzi de referinta", ["Rata de politică monetară BNR (%)", "Rata la 3 luni (%)"], {"zecimale": 2, "nume": {"Rata la 3 luni": "ROBOR 3 luni"}})]},
        ],
    },
    "asigurari": {
        "label": "Asigurări", "icon": "🛡️",
        "intro": "Cât s-au scumpit asigurările (RCA, locuință), câți bani încasează asigurătorii și cât plătesc daune, și cât de solizi sunt.",
        "seturi": [
            {"cod": "03 01", "key": "preturi", "scurt": "Prețul asigurărilor", "icon": "🏷️",
             "descriere": "Cât de scumpe sunt asigurările față de 2015 (indice: 100 = prețurile din 2015). 300 înseamnă de 3 ori mai scump. "
                          "Compară cu „Total coș de consum” = inflația generală.",
             "unitate": "indice, 2015 = 100", "zecimale": 1,
             "foi": [("Indici lunari RO", None, {})]},
            {"cod": "03 02", "key": "prime", "scurt": "Prime și daune", "icon": "🧾",
             "descriere": "Câți bani încasează asigurătorii din România din polițe (prime) și cât plătesc în daune și cheltuieli, pe trimestru, în milioane de euro.",
             "unitate": "mil. euro", "zecimale": 1,
             "foi": [("Prime trimestriale", ["Total prime brute (mil. EUR)", "Prime brute generale (mil. EUR)", "Prime brute viață (mil. EUR)"], {}),
                     ("Daune și cheltuieli", ["Daune brute generale (mil. EUR)", "Daune brute viață (mil. EUR)",
                                              "Cheltuieli de exploatare generale (mil. EUR)", "Cheltuieli de exploatare viață (mil. EUR)"], {})]},
            {"cod": "03 03", "key": "solvabilitate", "scurt": "Cât de solizi sunt asigurătorii", "icon": "🧱",
             "descriere": "Rata de solvabilitate arată de câte ori acoperă capitalul unui asigurător minimul cerut de lege. "
                          "Sub 100% e semnal de alarmă; 200% înseamnă de două ori minimul.",
             "unitate": "%", "zecimale": 1,
             "foi": [("Rata de solvabilitate", ["Rata de solvabilitate — total piață (%)", "Percentila P10 (%)", "Percentila P50 (%)", "Rata MCR — total piață (%)"],
                      {"nume": {"Percentila P10": "Cei mai slabi 10% dintre asigurători (P10)", "Percentila P50": "Asigurătorul median (P50)"}}),
                     ("Pe tipuri de societăți", None, {}),
                     ("Capital și cerințe", ["Fonduri proprii eligibile pentru SCR (mil. EUR)", "Cerința de capital de solvabilitate SCR (mil. EUR)", "Surplus de capital (mil. EUR)"], {})]},
        ],
    },
    "investitii": {
        "label": "Investiții financiare", "icon": "💼",
        "intro": "Unde își țin românii banii — depozite, fonduri de investiții, acțiuni — și câți bani intră sau ies din țară prin investiții.",
        "seturi": [
            {"cod": "04 01", "key": "fluxuri", "scurt": "Investiții cu străinătatea", "icon": "🌍",
             "descriere": "Câți bani au intrat sau au ieșit din România prin investiții, în fiecare lună, în milioane de euro. "
                          "Minus = au intrat mai mulți bani străini decât au investit românii în afară.",
             "unitate": "mil. euro", "zecimale": 1,
             "foi": [("Fluxuri nete lunare", None, {"suma": True})]},
            {"cod": "04 02", "key": "depozite", "scurt": "Depozitele românilor", "icon": "🐷",
             "descriere": "Câți bani țin în bănci populația, firmele și alte instituții, la sfârșitul fiecărei luni, în milioane de euro — și pe ce termen.",
             "unitate": "mil. euro", "zecimale": 0,
             "foi": [("Solduri pe sectoare", None, {}),
                     ("Depozitele populației", ["La vedere (mil. EUR)", "La termen sub 1 an (mil. EUR)", "La termen 1–2 ani (mil. EUR)", "La termen peste 2 ani (mil. EUR)"],
                      {"prefix": "Populație — "})]},
            {"cod": "04 03", "key": "fonduri", "scurt": "Fonduri de investiții", "icon": "🧺",
             "descriere": "Câți bani administrează fondurile de investiții din România, pe tipuri (acțiuni, obligațiuni, mixte, imobiliare), în milioane de euro.",
             "unitate": "mil. euro", "zecimale": 1,
             "foi": [("Active pe tipuri de fonduri", None, {"prefix": "Active — "}),
                     ("Activ net al investitorilor", ["Toate fondurile"], {"prefix": "Activul net al investitorilor — "}),
                     ("Bani noi în fonduri", ["Toate fondurile — flux net lunar (mil. EUR)"], {"nume": {"Toate fondurile — flux net lunar": "Bani noi intrați în fonduri (net)"}})]},
            {"cod": "04 04", "key": "avere", "scurt": "Averea financiară a familiilor", "icon": "👨‍👩‍👧",
             "descriere": "Ce au gospodăriile din România: numerar, depozite, acțiuni, fonduri, asigurări, pensii private — și cât datorează. "
                          "Milioane de lei, la sfârșitul fiecărui trimestru.",
             "unitate": "mil. lei", "zecimale": 0,
             "foi": [("Active financiare (lei)", ["Total active financiare", "Numerar și depozite", "Numerar (bani cash)", "Depozite la vedere (conturi curente)",
                                                  "Alte depozite (la termen, economii)", "Acțiuni și participații", "Acțiuni listate la bursă",
                                                  "Unități de fond de investiții", "Titluri de datorie (obligațiuni)", "Asigurări de viață",
                                                  "Drepturi de pensie (Pilon II și III)"], {}),
                     ("Avere financiară netă", ["Pasive financiare (mil. lei)", "Avere financiară netă (mil. lei)"],
                      {"nume": {"Pasive financiare": "Datorii (credite etc.)"}})]},
        ],
    },
    "pensii": {
        "label": "Pensii private", "icon": "👵",
        "intro": "Pilonul II (obligatoriu, din contribuția la pensie) și Pilonul III (facultativ): câți participanți, câți bani, unde sunt investiți și ce randament au.",
        "seturi": [
            {"cod": "05 01", "key": "pilon2", "scurt": "Pilonul II: participanți și bani", "icon": "👥",
             "descriere": "Câți români au cont la Pilonul II, câți bani s-au strâns în fonduri și cât se virează lunar din contribuții.",
             "unitate": "mil. lei", "zecimale": 1,
             "foi": [("Total sistem lunar", ["Participanți (mii pers.)", "Activ net (mil. lei)", "Activ mediu/participant (lei)",
                                             "Contribuții brute lunare (mil. lei)", "Contribuție medie/participant (lei)"], {"zecimale": 1}),
                     ("Activ net pe fond", None, {"prefix": "Activ net — "})]},
            {"cod": "05 02", "key": "plasamente", "scurt": "Pilonul II: unde sunt investiți banii", "icon": "🧮",
             "descriere": "Ce procent din banii de la Pilonul II e pus în titluri de stat, acțiuni, depozite, obligațiuni sau fonduri.",
             "unitate": "% din active", "zecimale": 2, "rata": True,
             "foi": [("Structura sistem (%)", None, {"fara": ["Activ total (mil. lei)"]})]},
            {"cod": "05 03", "key": "randamente", "scurt": "Pilonul II: randamente", "icon": "📈",
             "descriere": "VUAN = prețul unei unități de fond (a pornit de la 10 lei în 2008). Dacă a ajuns la 40, banii s-au înmulțit de 4 ori. "
                          "Randamentul anualizat ASF = cât a câștigat fondul pe an, în medie, în ultimii 5 ani.",
             "unitate": "lei / unitate", "zecimale": 4,
             "foi": [("VUAN lunar pe fond", None, {"prefix": "VUAN — "}),
                     ("Rate de rentabilitate ASF (%)", None, {"prefix": "Randament anualizat — ", "unitate": "% pe an", "zecimale": 2})]},
            {"cod": "05 04", "key": "pilon3", "scurt": "Pilonul III (facultativ)", "icon": "🌱",
             "descriere": "Pensiile private facultative: câți participanți, câți bani, cât se contribuie lunar și cât are fiecare fond.",
             "unitate": "mil. lei", "zecimale": 1,
             "foi": [("Total sistem lunar", ["Participanți (persoane)", "Activ net (mil. lei)", "Activ mediu/participant (lei)",
                                             "Contribuții brute lunare (lei)", "Contribuție medie/participant (lei)"], {"zecimale": 0}),
                     ("Activ net pe fond", None, {"prefix": "Activ net — ", "zecimale": 2})]},
        ],
    },
}


def ruleaza_scripturi() -> list[str]:
    erori = []
    for f in sorted((PKG / "scripts").glob("financiara_*.py")):
        print(f"── {f.name}", flush=True)
        try:
            r = subprocess.run([sys.executable, str(f)], cwd=PKG / "scripts", timeout=1800)
            if r.returncode:
                erori.append(f"{f.stem}: cod {r.returncode}")
        except subprocess.TimeoutExpired:
            erori.append(f"{f.stem}: timeout")
    return erori


def foaie(wb, nume: str) -> tuple[list[str], list[str], dict[str, list]]:
    """(antet, perioade, {coloană: valori}) din foaia cu rândul 3 = antet."""
    rows = list(wb[nume].iter_rows(min_row=3, values_only=True))
    antet = [str(c).strip() if c is not None else "" for c in rows[0]]
    rows = [r for r in rows[1:] if r[0] and re.match(r"^\d{4}-(\d{2}|Q\d)$", str(r[0]))]
    per = [str(r[0]) for r in rows]
    col = {a: [r[j] if isinstance(r[j], (int, float)) else None for r in rows] for j, a in enumerate(antet) if j and a}
    return antet, per, col


def sursa(wb) -> dict:
    """Câmpurile din foaia Sursa + lista „Note”."""
    d, note, sec = {}, [], None
    for a, b, *_ in wb["Sursa"].iter_rows(min_row=4, values_only=True):
        if a:
            sec = a
            if b and not str(b).startswith("•"):
                d[a] = str(b)
        if b and str(b).startswith("•") and sec == "Note":
            note.append(str(b).lstrip("• ").strip())
    return {**d, "note": note}


def rot(v, z):
    return None if v is None else (round(v, z) if z else int(round(v)))


def construieste(cfg: dict, cale: Path) -> dict:
    wb = openpyxl.load_workbook(cale, read_only=True, data_only=True)
    meta = sursa(wb)
    perioade: set[str] = set()
    brute = []
    for nume_foaie, coloane, opt in cfg["foi"]:
        antet, per, col = foaie(wb, nume_foaie)
        lipsa = [c for c in (coloane or []) if c not in col]
        if lipsa:
            raise RuntimeError(f"[{nume_foaie}] coloane lipsă: {lipsa}")
        for c in coloane or [a for a in antet[1:] if a and a not in opt.get("fara", [])]:
            m = RE_UNIT.search(c)
            nume = RE_UNIT.sub("", c).strip()
            nume = opt.get("prefix", "") + opt.get("nume", {}).get(nume, nume)
            unit = m.group(1) if m else opt.get("unitate", cfg.get("unitate", ""))
            unit = {"mil. EUR": "mil. euro", "mld. EUR": "mld. euro", "mil. RON": "mil. lei", "mld. RON": "mld. lei",
                    "RON": "lei", "mii pers.": "mii persoane", "p.p.": "puncte procentuale"}.get(unit, unit)
            brute.append({"nume": nume, "valori": dict(zip(per, col[c])), "unitate": unit,
                          "zecimale": opt.get("zecimale", cfg["zecimale"]),
                          "agregare": "suma" if opt.get("suma") or RE_SUMA.search(c) else "medie"})
        perioade.update(per)
    # același nume în unități diferite (ex. capitalizarea în lei și în euro) → adăugăm unitatea
    for s in brute:
        if sum(x["nume"] == s["nume"] for x in brute) > 1:
            s["nume_u"] = f'{s["nume"]} ({s["unitate"]})'
    for s in brute:
        s["nume"] = s.pop("nume_u", s["nume"])
    # tăiem perioadele goale de la început și de la sfârșit
    per = sorted(perioade)
    plin = [p for p in per if any(s["valori"].get(p) is not None for s in brute)]
    per = per[per.index(plin[0]):per.index(plin[-1]) + 1]
    unit_set = cfg.get("unitate") or brute[0]["unitate"]
    serii = []
    for s in brute:
        r = {"nume": s["nume"], "valori": [rot(s["valori"].get(p), s["zecimale"]) for p in per], "provizorii": []}
        if s["unitate"] != unit_set:
            r["unitate"] = s["unitate"]
        if s["zecimale"] != cfg["zecimale"]:
            r["zecimale"] = s["zecimale"]
        if s["agregare"] != "medie":
            r["agregare"] = s["agregare"]
        if any(v is not None for v in r["valori"]):
            serii.append(r)
    url = meta.get("URL sursă", "")
    return {
        "key": cfg["key"], "scurt": cfg["scurt"], "icon": cfg["icon"], "titlu": TITLURI.get(cfg["cod"], cfg["scurt"]),
        "descriere": cfg["descriere"], "cod": re.split(r"[;,(]| — |: ", meta.get("Cod / identificator set", ""))[0].strip()[:40],
        "freq": "Q" if "Q" in per[-1] else "M", "unitate": unit_set, "zecimale": cfg["zecimale"], "agregare": "medie",
        "rata": cfg.get("rata", False), "sursa": meta.get("Sursa datelor", "").split(" — ")[0].split(" (")[0],
        "note": meta["note"][:5], "url": url if url.startswith("http") else "https://24reco.com",
        "actualizat": date.today().isoformat(), "perioade": per, "serii": serii,
    }


def titluri() -> dict:
    """Titlurile lungi din foaia Sursa (rândul 1) ale fiecărui Excel."""
    t = {}
    for f in XLSX_DIR.glob("Financiara ?? ??.xlsx"):
        try:
            t[f.stem[-5:]] = openpyxl.load_workbook(f, read_only=True)["Sursa"]["A1"].value
        except Exception:  # noqa: BLE001
            pass
    return t


TITLURI: dict = {}


def main():
    global XLSX_DIR, TITLURI
    erori = []
    if "--doar-conversie" in sys.argv:
        rest = [a for a in sys.argv[1:] if not a.startswith("--")]
        if rest:
            XLSX_DIR = Path(rest[0])
    else:
        erori += ruleaza_scripturi()
    TITLURI = titluri()
    vechi = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else {}
    rezultat = {}
    for tkey, tema in SETURI.items():
        v_tema = {s["key"]: s for s in vechi.get(tkey, {}).get("seturi", [])}
        lista = []
        for cfg in tema["seturi"]:
            cale = XLSX_DIR / f"Financiara {cfg['cod']}.xlsx"
            try:
                if not cale.exists():
                    raise RuntimeError(f"lipsește {cale.name}")
                s = construieste(cfg, cale)
                v = v_tema.get(cfg["key"])
                if v and v["perioade"] == s["perioade"] and v["serii"] == s["serii"]:
                    s["actualizat"] = v["actualizat"]  # date neschimbate → fără commit inutil
                print(f"[{cfg['cod']}] ok: {len(s['perioade'])} perioade {s['perioade'][0]} → {s['perioade'][-1]}, {len(s['serii'])} serii")
                lista.append(s)
            except Exception as e:  # păstrăm datele vechi ale setului
                erori.append(f"{cfg['cod']}: {e}")
                print(f"[{cfg['cod']}] EROARE: {e}")
                if cfg["key"] in v_tema:
                    lista.append(v_tema[cfg["key"]])
        rezultat[tkey] = {"label": tema["label"], "icon": tema["icon"], "intro": tema["intro"], "seturi": lista}
    if not any(t["seturi"] for t in rezultat.values()):
        sys.exit("::error::Niciun set nu a putut fi generat.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(rezultat, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Scris {OUT} ({OUT.stat().st_size // 1024} KB)")
    if erori:
        print("::warning::Seturi nerefăcute (s-au păstrat datele vechi): " + "; ".join(erori))


if __name__ == "__main__":
    main()
