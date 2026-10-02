#!/usr/bin/env python3
"""
Prețul motorinei în UE — Weekly Oil Bulletin (Comisia Europeană), săptămânal.

Descarcă „Price developments 2005 onwards (xlsx)” de pe pagina Weekly Oil Bulletin
și rescrie `dieselData` din public/transport-data.json (EUR/litru, cu taxe, per țară + media UE).
Linkul fișierului se schimbă la fiecare publicare, așa că îl căutăm în pagină.

Rulare: pip install openpyxl requests ; python scripts/fetch_oil_bulletin.py
"""
import io
import json
import re
import sys
from pathlib import Path

import openpyxl
import requests

PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
BASE = "https://energy.ec.europa.eu"
DATA_FILE = Path("public/transport-data.json")
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120 (+https://24reco.com)"}


def history_url() -> str:
    html = requests.get(PAGE, headers=HEADERS, timeout=60).text
    links = re.findall(r'href="(/document/download/[^"]+?\.xlsx)"', html)
    hist = [l for l in links if "History" in l or "history" in l]
    if not hist:
        sys.exit("::error::Nu găsesc fișierul „Prices History” pe pagina Weekly Oil Bulletin — s-a schimbat pagina?")
    return BASE + hist[0].replace("&amp;", "&")


def main():
    url = history_url()
    print("Descarc", url)
    r = requests.get(url, headers=HEADERS, timeout=180)
    r.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["Prices with taxes"]
    header = next(ws.iter_rows(max_row=1, values_only=True))
    cols = {}
    for i, h in enumerate(header):
        m = re.fullmatch(r"([A-Z]{2,3})_price_with_tax_diesel", str(h or ""))
        if m and m.group(1) != "EUR":  # EUR = zona euro; păstrăm doar UE + țările
            cols[m.group(1)] = i

    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            break
        d = {"date": row[0].strftime("%Y-%m-%d")}
        for cc, i in cols.items():
            v = row[i]
            d[cc] = round(v / 1000, 4) if isinstance(v, (int, float)) and v > 0 else None  # €/1000 l → €/l
        rows.append(d)
    rows.sort(key=lambda x: x["date"], reverse=True)
    if len(rows) < 500 or "RO" not in cols:
        sys.exit(f"::error::Date suspecte: {len(rows)} săptămâni, coloane {sorted(cols)}")

    data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
    vechi = data.get("dieselData", [])
    data["dieselData"] = rows
    data.setdefault("countryNames", {})["EU"] = "Media UE"
    DATA_FILE.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"dieselData: {len(vechi)} → {len(rows)} săptămâni, ultima {rows[0]['date']} "
          f"(RO {rows[0]['RO']} €/l, UE {rows[0]['EU']} €/l)")


if __name__ == "__main__":
    main()
